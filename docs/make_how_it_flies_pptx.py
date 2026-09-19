"""Build docs/how_it_flies.pptx from the walkthrough (python-pptx). Run from docs/."""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)


def slide(title, bullets, sub=None):
    s = prs.slides.add_slide(prs.slide_layouts[5])          # title only
    s.shapes.title.text = title
    s.shapes.title.text_frame.paragraphs[0].font.size = Pt(34)
    box = s.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(12.0), Inches(5.5))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for b in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        level = 1 if b.startswith("  ") else 0
        p.text = ("- " if level == 0 else "   - ") + b.strip()
        p.font.size = Pt(22 if level == 0 else 18)
        p.space_after = Pt(8)
    if sub:
        p = tf.add_paragraph(); p.text = sub; p.font.size = Pt(16); p.font.italic = True
    return s


s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "How the drone flies the course"
s.placeholders[1].text = "AI-GrandPrix, physical qualifier. Start to finish, no prior knowledge assumed."

slide("The one-minute version", [
    "The organizers published where every gate stands. We trust that map.",
    "Before the race a planner draws the fastest line through the gates that our camera and drone can handle.",
    "In the air there is no GPS. The drone counts its steps with the flight controller's motion sensors: eyes closed.",
    "When the camera sees the next gate where the map says it must be, it opens its eyes and corrects.",
    "A follower steers along the planned line and sends stick commands 50 times a second.",
    "A human on the radio arms it and can take over at any moment.",
])

s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = "The whole loop on one page"
s.shapes.add_picture(str(HERE / "how_it_flies.png"), Inches(1.6), Inches(1.2), height=Inches(6.1))

slide("1. The map", [
    "data/course_map.json: the organizer's table. 10 gates, position, crossing direction, height.",
    "Gate 9 is a double: through the top at 4 m, U-turn, back through the bottom at 1.35 m.",
    "Two laps = 23 crossings. The start line is 7 m behind gate 1.",
    "Nothing is surveyed at the venue. The map is the truth we plan on.",
])

slide("2. The levers", [
    "config/vehicle_cam35_120.toml holds everything the planner needs to know about the drone:",
    "  the camera: mount angle and lens field of view, measured on the drone with hardware.camcal",
    "  the limits: max tilt, how fast it may roll into a turn, lateral margin, top speed, climb rate",
    "One number k moves all limits together: k = 0 is the safe floor, k = 1 the race values.",
    "No target time anywhere. The time is whatever the solve gives.",
])

slide("3. The planner", [
    "python -m raceline.ladder --k 0.5 --config ../config/vehicle_cam35_120.toml",
    "Draws a smooth line through the gate centres, then a speed profile inside the limits.",
    "Refuses to lean the drone so far that the camera would lose the next gate on the approach.",
    "Writes the plan: out/plans/plan_LADDER_k050.json, the line, the speed at every point, the 23 crossings.",
    "Prints the model time. k = 0.33 gives about 56 s, k = 0.5 about 49 s, k = 1 about 30 s.",
])

slide("4. What the flight controller gives us", [
    "Betaflight 4.4.3 on the Archer, talked to over a serial link (MSP), about 50 times a second:",
    "  attitude (how tilted), heading (which way the nose points)",
    "  accelerometer (how hard it is pushed), barometer (altitude)",
    "No GPS, no position, no magnetometer. Heading is gyro-integrated and zeroed on the start line.",
    "Measured from the organizers' own flight logs: accel scale 2048 counts per g, baro noise 2.5 cm.",
])

slide("5. Knowing where you are: dead reckoning", [
    "Rotate the accelerometer into the world with the attitude, subtract gravity, integrate twice.",
    "Velocity and position come out. They drift: a few tenths of a metre every few seconds.",
    "Altitude: accelerometer and barometer blended. Each alone is bad, together they are good.",
    "Run on the organizers' real 82 s flight log this drifts 0.14 m/s^2, which the sim tolerates.",
])

slide("6. Knowing where you are: the camera", [
    "The detector finds red blobs. It does not know which blob is which gate.",
    "The estimator knows, from the map and its own position, exactly where the next gate must appear.",
    "A blob at that azimuth with the right apparent size is the gate. Its bearing and size give a position fix.",
    "Only the gate just passed and the next one may ever match. Everything else is ignored.",
    "Range comes from the blob width. The real gate has a header board, so the width, not the height.",
])

slide("7. Knowing where you are: the crossing itself", [
    "When the estimate crosses a gate plane, the estimator counts the crossing itself. No referee on the drone.",
    "At that instant the drone was inside a 1.5 m opening, so the lateral estimate is pulled to the gate centre.",
    "Every blind turn therefore starts from a known point.",
    "Through the hairpin and the stacked gate the next gate is out of view for 1 to 2 s; the follower flies the known turn from the plan.",
])

slide("8. Steering: the follower", [
    "Where on the line am I. Aim at a carrot a little ahead. Acceleration to get there, plus the plan's own acceleration.",
    "That becomes a tilt angle (ANGLE mode: the flight controller levels itself), a throttle for altitude, a yaw for the nose.",
    "The nose is pointed at the next gate through the turns so the camera looks where the gate must be.",
    "Sticks go over MSP. The flight controller does the fast attitude loop itself.",
])

slide("9. The human", [
    "MSP override covers the four sticks only (msp_override_channels_mask = 15, set once in the CLI).",
    "The pilot arms, flips MSP OVERRIDE and ANGLE on the radio.",
    "The runtime waits for both before the plan clock starts.",
    "The pilot can flip override off and have the sticks back, or disarm, at any moment.",
])

slide("10. Race day", [
    "Bench: bench info, rc-test, arm-test, drift; camcal on a real gate; dry run on the start line.",
    "Fly the lowest k that was clean in the sim. Clean twice, jump to the fastest k brought. Fail, fly the midpoint.",
    "A complete slow run scores above any incomplete fast one.",
])

slide("11. Proven and not proven", [
    "Sim, noisy synthetic detector, no ground truth: k 0.4 to 0.5 completes about 2 runs in 3 at 47 to 50 s.",
    "Slower rungs complete nearly always. The 30 s plan does not complete on vision.",
    "Never flown on the real drone. The follower gains have never met the real airframe.",
    "The cage flight answers what the sim cannot.",
])

out = HERE / "how_it_flies.pptx"
prs.save(str(out))
print("wrote", out, len(prs.slides), "slides")
